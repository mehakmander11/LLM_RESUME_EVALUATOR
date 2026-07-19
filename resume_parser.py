import os 
from pathlib import Path
from dotenv import load_dotenv
from groq import Groq
from pydantic import BaseModel
import time

load_dotenv()

my_api_key = os.getenv("GROQ_API_KEY")

if not my_api_key:
    raise ValueError("Not found your api key!")

client = Groq(api_key=my_api_key)
model = "llama-3.3-70b-versatile"

#Paste job decription here
job_description = """
We're looking for exceptional students who love building software and solving difficult engineering problems.

Education
Pursuing a Bachelor's or Integrated Master's degree
Expected graduation in 2028
Mathematics, Computer Science or Electrical Engineering
CGPA of 8.5 or above
No active backlogs
Technical Foundation
You should have:

Strong programming skills in C++ and/or Python.
Excellent understanding of data structures, algorithms and object-oriented programming.
Familiarity with operating systems, computer architecture and networking fundamentals.
Strong analytical thinking and problem-solving ability.
Passion for writing clean, efficient and reliable software.
What We Value: Beyond technical skills, we're interested in how you think.
We look for people who are:

Curious enough to question existing solutions.
Analytical enough to simplify complex systems.
Persistent enough to debug difficult problems.
Detail-oriented enough to care about every line of code.
Collaborative enough to learn from others and improve together.
Motivated by solving engineering challenges that have real-world impact
"""

class jobD(BaseModel):
    role:str
    required_skils:list[str]
    preferred_skils:list[str]
    minimum_experience:float | None
    education_requirements:list[str]
    responsibilities:list[str]

jobd_schema = jobD.model_json_schema()

sytem_prompt = f"""
You are an expert HR assistant.

Your job is to analyze the job description and extract
the following information from it.

Return only valid JSON matching the following schema:
{jobd_schema}

IMPORTANT:
DO NOT return schema itself.
DO NOT return fields like "properties", "title", or "type".
Fill schema fields with the information extracted from the job description.

If minimum_experience is not mentioned in the job description, return null for that field.
If information for list fields is not mentioned in the job description, return an empty list for that field.
Do not invent any information.
"""

user_prompt = f"""
Analysze the following job description
{job_description}
"""

message_system = {"role": "system", "content": sytem_prompt}
message_user = {"role": "user", "content": user_prompt}

response_format = {"type" : "json_object"}

messages = [message_system, message_user]

response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)

ans = response.choices[0].message.content
raw_json = ans

#convert text answer to json object
import json
job_data = json.loads(raw_json)

job = jobD(**job_data)

print(job.minimum_experience)
print(job.role)

class MatchResult(BaseModel):
    score: float
    details: dict
class Experience(BaseModel):
    company: str | None = None
    role: str | None = None
    duration: str | None = None
    description: str | None = None
    skills_used: list[str] = []

class Resume(BaseModel):
    name: str | None = None
    email: str | None = None
    phone: str | None = None

    total_experience_years: float | None = None

    skills: list[str] = []
    experiences: list[Experience] = []
    education: list[str] = []
    projects: list[str] = []
    certifications: list[str] = []


resume_schema = Resume.model_json_schema()
def final_score(job,resume):
    match_schema = MatchResult.model_json_schema()
    prompt = f"""
    You are an HR recruiter.

    Compare the candidate's resume with the job description.

    JOB DESCRIPTION:
    {job.model_dump_json(indent=2)}

    CANDIDATE RESUME:
    {resume.model_dump_json(indent=2)}
    Return JSON matching this schema:

    {match_schema}

    Give me:

    1. Candidate name
    2. Matching skills
    3. Missing important skills
    4. Whether experience requirement is met
    5. Overall match percentage from 0 to 100
    6. A short final verdict

    Keep the response concise and easy to read.
    """
    message={
        "role": "user",
        "content" : prompt
    }
    messages=[message]
    response_format={
        "type": "json_object"
    }
    response = client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    data = json.loads(response.choices[0].message.content)
    return MatchResult(**data)
def parse_resume(resume_text):
    system_prompt = f"""
    You are an expert resume parser.

    Extract information from the resume based on its meaning,
    not only based on exact section headings.

    Different resumes may use different headings.

    For example:
    - Experience
    - Professional Experience
    - Work History
    - Employment
    - Internships

    These may all contain relevant experience.

    Skills may also appear in the skills section, work experience,
    internships or projects.

    Return ONLY valid JSON matching this schema:

    {resume_schema}

    Important rules:

    1. Do not invent information.
    2. If a value is not available, return null.
    3. If a list has no information, return an empty list.
    4. Include internships inside experiences.
    5. Extract skills mentioned across the entire resume.
    """
    user_prompt = f"""
    Parse the following resume:

    {resume_text}
    """
    message_system={
        "role" : "system",
        "content" : system_prompt
    }
    message_user={
        "role" : "user",
        "content" : user_prompt
    }
    messages=[message_system, message_user]
    response_format={
        "type": "json_object"
    }
    response=client.chat.completions.create(model=model, messages=messages, response_format=response_format)
    raw_output = response.choices[0].message.content
    data = json.loads(raw_output)
    resume = Resume(**data)
    return resume


from pypdf import PdfReader
from docx import Document
def read_pdf(file_path):
    reader = PdfReader(file_path)
    text = ""
    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + "\n"
    return text

def read_docx(file_path):
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    text += cell.text + "\n"
    return text


def read_resume(file_path):
    if file_path.suffix.lower() == ".pdf":
        return read_pdf(file_path)
    elif file_path.suffix.lower() == ".docx":
        return read_docx(file_path)
    else:
        return None



# lets do it now
resume_folder = Path("resumes")
all_results=[]
for file_path in resume_folder.iterdir():
    #C:\Users\Pratyush\padho_with_pratyush\week1\day5\resumes\abhay resume new - Abhay Singh.pdf
    if file_path.suffix.lower() not in [".pdf", ".docx"]:
        continue
    print("\nProcessing:", file_path.name)
    resume_text = read_resume(file_path)
    parsed_resume=parse_resume(resume_text) # llm call1
    time.sleep(5)
    result = final_score(job, parsed_resume) #llm caLL2
    #score and details
    #acount chtgpt
    # request bhejna shhur krega millions
    #chattgot server jam ho jayega
    time.sleep(5)
    print("Score:", result.score)
    all_results.append({
        "name": parsed_resume.name,
        "score": result.score,
        "details": result.details
    })
all_results.sort(
    key=lambda candidate: candidate["score"],
    reverse=True
)
top_2 = all_results[:2]
worst_2 = all_results[-2:]


print("TOP 2 CANDIDATES")
for candidate in top_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )

    print(candidate["details"])

print("LOWEST 2 CANDIDATES")
for candidate in worst_2:

    print(
        candidate["name"],
        "-",
        candidate["score"],
        "%"
    )
    print(candidate["details"])